"""Basic AI Knowledge Assistant - Lite

Supports common text-based knowledge sources:
PDF, PowerPoint, Word, TXT, Markdown, CSV, Excel, JSON,
Google Sheets, and public web pages.

Architecture:
Sources -> text extraction -> chunking -> FastEmbed -> Qdrant -> RAG -> Groq LLM
                                                         ^
                                                         |
                                                  chat memory
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
from typing import Iterable

import requests
import streamlit as st
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from qdrant_client import QdrantClient, models
from langchain_core.chat_history import InMemoryChatMessageHistory
from langchain_core.messages import AIMessage, HumanMessage
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_groq import ChatGroq
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation

load_dotenv()

MODEL_NAME = os.getenv("EMBEDDING_MODEL", "BAAI/bge-small-en-v1.5")
GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")
TOP_K = int(os.getenv("TOP_K", "4"))
CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "900"))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "150"))
COLLECTION = "knowledge"

SUPPORTED = ["pdf", "pptx", "docx", "txt", "md", "csv", "xlsx", "json"]


# ------------------------- Source extraction -------------------------

def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_pdf(data: bytes, name: str) -> list[dict]:
    reader = PdfReader(io.BytesIO(data))
    out = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            out.append({"text": text, "source": name, "location": f"page {page_no}"})
    return out


def extract_pptx(data: bytes, name: str) -> list[dict]:
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        pieces = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                pieces.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    pieces.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = clean_text("\n".join(pieces))
        if text:
            out.append({"text": text, "source": name, "location": f"slide {slide_no}"})
    return out


def extract_docx(data: bytes, name: str) -> list[dict]:
    doc = DocxDocument(io.BytesIO(data))
    pieces = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            pieces.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = clean_text("\n".join(pieces))
    return [{"text": text, "source": name, "location": "document"}] if text else []


def extract_plain(data: bytes, name: str) -> list[dict]:
    text = clean_text(data.decode("utf-8", errors="ignore"))
    return [{"text": text, "source": name, "location": "text file"}] if text else []


def extract_csv(data: bytes, name: str) -> list[dict]:
    text = data.decode("utf-8-sig", errors="ignore")
    rows = csv.reader(io.StringIO(text))
    rows = list(rows)
    if not rows:
        return []
    headers = rows[0]
    out = []
    for i, row in enumerate(rows[1:], start=2):
        pairs = []
        for j, value in enumerate(row):
            if j < len(headers) and value.strip():
                pairs.append(f"{headers[j]}: {value.strip()}")
        if pairs:
            out.append({"text": " | ".join(pairs), "source": name, "location": f"row {i}"})
    return out


def extract_xlsx(data: bytes, name: str) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(x).strip() if x is not None else f"Column {i+1}" for i, x in enumerate(rows[0])]
        for row_no, row in enumerate(rows[1:], start=2):
            pairs = []
            for j, value in enumerate(row):
                if value not in (None, ""):
                    col = headers[j] if j < len(headers) else f"Column {j+1}"
                    pairs.append(f"{col}: {value}")
            if pairs:
                out.append({
                    "text": " | ".join(pairs),
                    "source": name,
                    "location": f"sheet {ws.title}, row {row_no}",
                })
    return out


def extract_json(data: bytes, name: str) -> list[dict]:
    obj = json.loads(data.decode("utf-8", errors="ignore"))
    text = clean_text(json.dumps(obj, ensure_ascii=False, indent=2))
    return [{"text": text, "source": name, "location": "JSON document"}] if text else []


def extract_file(uploaded_file) -> list[dict]:
    name = uploaded_file.name
    ext = name.lower().split(".")[-1]
    data = uploaded_file.getvalue()

    if ext == "pdf":
        return extract_pdf(data, name)
    if ext == "pptx":
        return extract_pptx(data, name)
    if ext == "docx":
        return extract_docx(data, name)
    if ext in {"txt", "md"}:
        return extract_plain(data, name)
    if ext == "csv":
        return extract_csv(data, name)
    if ext == "xlsx":
        return extract_xlsx(data, name)
    if ext == "json":
        return extract_json(data, name)

    raise ValueError(f"Unsupported file type: .{ext}")


def google_sheet_to_records(url: str) -> list[dict]:
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", url)
    if not match:
        raise ValueError("Invalid Google Sheets URL.")
    sheet_id = match.group(1)
    gid_match = re.search(r"(?:[#?&]gid=)(\d+)", url)
    gid = gid_match.group(1) if gid_match else "0"
    csv_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"
    response = requests.get(csv_url, timeout=20)
    response.raise_for_status()
    class Obj:
        def __init__(self, payload: bytes, name: str):
            self._payload = payload
            self.name = name
        def getvalue(self):
            return self._payload
    return extract_csv(response.content, "Google Sheet")


def web_page_to_records(url: str) -> list[dict]:
    if not url.startswith(("http://", "https://")):
        raise ValueError("URL must start with http:// or https://")
    response = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = clean_text(soup.get_text(" "))
    return [{"text": text, "source": url, "location": "web page"}] if text else []


# ------------------------- Chunking + vector store -------------------------

def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = text.strip()
    if len(text) <= size:
        return [text] if text else []
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        chunk = text[start:end]
        # Try to end at a sentence/space for cleaner chunks.
        if end < len(text):
            cut = max(chunk.rfind(". "), chunk.rfind(" "))
            if cut > size * 0.55:
                end = start + cut + 1
                chunk = text[start:end]
        chunks.append(chunk.strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def make_chunks(records: Iterable[dict]) -> list[dict]:
    chunks = []
    for record in records:
        for i, part in enumerate(chunk_text(record["text"]), start=1):
            chunks.append({
                "text": part,
                "source": record["source"],
                "location": record["location"],
                "chunk": i,
            })
    return chunks


def build_vector_store(chunks: list[dict]) -> QdrantClient:
    client = QdrantClient(":memory:")
    client.create_collection(
        collection_name=COLLECTION,
        vectors_config=models.VectorParams(
            size=client.get_embedding_size(MODEL_NAME),
            distance=models.Distance.COSINE,
        ),
    )
    payload = [
        {"document": c["text"], "source": c["source"], "location": c["location"], "chunk": c["chunk"]}
        for c in chunks
    ]
    vectors = [models.Document(text=c["text"], model=MODEL_NAME) for c in chunks]
    ids = list(range(1, len(chunks) + 1))
    client.upload_collection(collection_name=COLLECTION, vectors=vectors, payload=payload, ids=ids)
    return client


def retrieve(client: QdrantClient, query: str, top_k: int = TOP_K) -> list[dict]:
    hits = client.query_points(
        collection_name=COLLECTION,
        query=models.Document(text=query, model=MODEL_NAME),
        limit=top_k,
        with_payload=True,
    ).points
    return [hit.payload for hit in hits if hit.payload]


# ------------------------- LLM + memory -------------------------

def get_llm() -> ChatGroq:
    key = os.getenv("GROQ_API_KEY", "").strip()
    if not key:
        raise RuntimeError("GROQ_API_KEY is missing in your .env file.")
    return ChatGroq(model=GROQ_MODEL, temperature=0, groq_api_key=key)


def history_text(history: InMemoryChatMessageHistory) -> str:
    lines = []
    for m in history.messages:
        role = "User" if isinstance(m, HumanMessage) else "Assistant"
        lines.append(f"{role}: {m.content}")
    return "\n".join(lines)


def rewrite_question(question: str, history: InMemoryChatMessageHistory, llm: ChatGroq) -> str:
    if not history.messages:
        return question
    prompt = ChatPromptTemplate.from_messages([
        (
            "system",
            "Rewrite the latest question as one standalone search query. "
            "Use the conversation only to resolve references such as it, that, they, or the previous item. "
            "Return only the search query.",
        ),
        (
            "human",
            "Conversation:\n{history}\n\nLatest question:\n{question}",
        ),
    ])
    result = llm.invoke(prompt.format_messages(history=history_text(history), question=question))
    return str(result.content).strip()


def answer_question(question: str, client: QdrantClient, history: InMemoryChatMessageHistory, llm: ChatGroq):
    standalone = rewrite_question(question, history, llm)
    docs = retrieve(client, standalone)
    context = "\n\n".join(
        f"SOURCE: {d.get('source')} ({d.get('location')})\n{d.get('document', '')}" for d in docs
    )

    prompt = ChatPromptTemplate.from_messages([
        (
            "system",
            "You are a knowledge assistant. Answer using the supplied knowledge whenever possible. "
            "Do not invent facts. If the knowledge does not contain the answer, say that it is not available "
            "in the current knowledge base. Use previous conversation to understand follow-up questions. "
            "Keep the answer clear and concise.\n\nKNOWLEDGE:\n{context}",
        ),
        MessagesPlaceholder(variable_name="history"),
        ("human", "{question}"),
    ])

    messages = prompt.format_messages(
        context=context or "No relevant knowledge was retrieved.",
        history=history.messages,
        question=question,
    )
    result = llm.invoke(messages)
    answer = str(result.content)

    history.add_user_message(question)
    history.add_ai_message(answer)

    return answer, docs, standalone


# ------------------------- Streamlit UI -------------------------

st.set_page_config(page_title="AI Knowledge Assistant Lite", page_icon="🤖")
st.title("🤖 AI Knowledge Assistant — Lite")
st.caption("Upload your knowledge, ask questions, and test conversational memory.")

if "history" not in st.session_state:
    st.session_state.history = InMemoryChatMessageHistory()
if "vector_store" not in st.session_state:
    st.session_state.vector_store = None
if "source_names" not in st.session_state:
    st.session_state.source_names = []
if "messages" not in st.session_state:
    st.session_state.messages = []

with st.sidebar:
    st.header("1. Add knowledge")
    files = st.file_uploader(
        "Upload files",
        accept_multiple_files=True,
        help="Supported: PDF, PPTX, DOCX, TXT, MD, CSV, XLSX, JSON",
    )
    sheet_url = st.text_input("Google Sheet URL (optional)")
    web_url = st.text_input("Web page URL (optional)")

    if st.button("Build Knowledge Base", use_container_width=True):
        try:
            records = []
            names = []
            with st.spinner("Reading sources and creating the vector index..."):
                for uploaded in files or []:
                    records.extend(extract_file(uploaded))
                    names.append(uploaded.name)
                if sheet_url.strip():
                    records.extend(google_sheet_to_records(sheet_url.strip()))
                    names.append("Google Sheet")
                if web_url.strip():
                    records.extend(web_page_to_records(web_url.strip()))
                    names.append(web_url.strip())

                if not records:
                    st.error("Add at least one file, Google Sheet, or web URL.")
                else:
                    chunks = make_chunks(records)
                    st.session_state.vector_store = build_vector_store(chunks)
                    st.session_state.source_names = names
                    st.session_state.history = InMemoryChatMessageHistory()
                    st.session_state.messages = []
                    st.success(f"Loaded {len(names)} source(s) → {len(chunks)} chunks.")
        except Exception as exc:
            st.error(f"Could not build the knowledge base: {exc}")

    if st.button("Clear chat", use_container_width=True):
        st.session_state.history = InMemoryChatMessageHistory()
        st.session_state.messages = []
        st.rerun()

    if st.session_state.source_names:
        st.divider()
        st.write("**Loaded sources**")
        for name in st.session_state.source_names:
            st.write(f"• {name}")

    st.divider()
    st.write("**Pipeline**")
    st.code("Source → Chunk → Embedding → Qdrant → RAG → Groq → Memory")

for msg in st.session_state.messages:
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])
        if msg.get("sources"):
            with st.expander("Sources used"):
                for source in msg["sources"]:
                    st.write(f"- {source}")

question = st.chat_input("Ask something about your knowledge base...")

if question:
    if st.session_state.vector_store is None:
        st.warning("Build the knowledge base first from the sidebar.")
    else:
        with st.chat_message("user"):
            st.markdown(question)
        try:
            with st.spinner("Thinking..."):
                llm = get_llm()
                answer, docs, standalone = answer_question(
                    question,
                    st.session_state.vector_store,
                    st.session_state.history,
                    llm,
                )
            source_labels = [f"{d.get('source')} — {d.get('location')}" for d in docs]
            st.session_state.messages.append({"role": "user", "content": question})
            st.session_state.messages.append({"role": "assistant", "content": answer, "sources": source_labels})
            with st.chat_message("assistant"):
                st.markdown(answer)
                with st.expander("Retrieval details"):
                    st.write(f"Search query: {standalone}")
                    for label in source_labels:
                        st.write(f"- {label}")
        except Exception as exc:
            st.error(str(exc))
